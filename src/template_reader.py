"""
template_reader.py
==================
Extracts the 248 checkpoints from the SEPG Master Audit Template Excel file
and writes them to checkpoint_registry.json.

Column mapping (sheet "Q2 2026"):
  A = Sr. No.       (phase number  — sparse, forward-filled)
  B = Phases/Process (phase name   — sparse, forward-filled)
  C = Activities    (activity name — sparse, forward-filled)
  D = Checkpoint    (merged D:K   — read from D top-left)
  L = Response      (verdict cell — where we WRITE to)
  M = Weighted      (weight       — integer 1-5)
  Q = observation   (evidence     — where we WRITE to)

Columns N, O, P, G9, X are formula cells — NEVER written to.
Rows where D is empty are boundary/separator rows — skipped.
"""

import json
import logging
import re
from pathlib import Path

import openpyxl
import yaml

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s — %(message)s",
)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _col_letter_to_index(letter: str) -> int:
    """Convert a column letter (e.g. 'A') to a 1-based column index."""
    return openpyxl.utils.column_index_from_string(letter)


def _cell_value(ws, row: int, col_letter: str):
    """Return the value of a cell, handling merged-cell top-left resolution."""
    col_idx = _col_letter_to_index(col_letter)
    cell = ws.cell(row=row, column=col_idx)
    # If the cell is a MergedCell openpyxl returns None — resolve to top-left
    if cell.value is None:
        # Check whether this cell is inside a merged range
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                # top-left cell of the merge
                tl_row = merged_range.min_row
                tl_col = merged_range.min_col
                return ws.cell(row=tl_row, column=tl_col).value
    return cell.value


def _to_float_weight(raw) -> float:
    """Convert raw weight value to float; return 0.0 if conversion fails."""
    if raw is None:
        return 0.0
    try:
        return float(raw)
    except (ValueError, TypeError):
        return 0.0


# ---------------------------------------------------------------------------
# Main extraction function
# ---------------------------------------------------------------------------

def extract_checkpoints(config: dict, output_path: str = "checkpoint_registry.json") -> list[dict]:
    """
    Read the master template and return the list of 248 checkpoint dicts.

    Args:
        config:      Loaded config.yaml as a dict.
        output_path: Where to write checkpoint_registry.json.

    Returns:
        List of checkpoint dicts.

    Raises:
        AssertionError: If the count != expected or IDs are not unique.
    """
    tmpl_cfg = config["template"]
    path = tmpl_cfg["path"]
    sheet_name = tmpl_cfg["sheet_name"]
    cols = tmpl_cfg["columns"]

    start_row = tmpl_cfg["data_start_row"]   # 12
    end_row   = tmpl_cfg["data_end_row"]     # 299
    expected  = tmpl_cfg["expected_checkpoint_count"]  # 248

    col_a = cols["sr_no"]
    col_b = cols["phase_name"]
    col_c = cols["activity"]
    col_d = cols["checkpoint_text"]
    col_l = cols["response_verdict"]
    col_m = cols["weight"]
    col_q = cols["observation"]

    logger.info("Loading workbook: %s", path)
    wb = openpyxl.load_workbook(path, data_only=False)

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")

    ws = wb[sheet_name]
    logger.info("Reading sheet '%s' (rows %d–%d)", sheet_name, start_row, end_row)

    checkpoints: list[dict] = []

    # Forward-fill state
    last_phase_number = ""
    last_phase_name   = ""
    last_activity     = ""

    for row in range(start_row, end_row + 1):
        # --- Forward-fill phase number (col A) ---
        raw_a = _cell_value(ws, row, col_a)
        if raw_a is not None and str(raw_a).strip():
            last_phase_number = str(raw_a).strip()

        # --- Forward-fill phase name (col B) ---
        raw_b = _cell_value(ws, row, col_b)
        if raw_b is not None and str(raw_b).strip():
            last_phase_name = str(raw_b).strip()

        # --- Forward-fill activity (col C) ---
        raw_c = _cell_value(ws, row, col_c)
        if raw_c is not None and str(raw_c).strip():
            last_activity = str(raw_c).strip()

        # --- Checkpoint text (col D, merged D:K) ---
        raw_d = _cell_value(ws, row, col_d)
        if raw_d is None:
            continue  # boundary row — no checkpoint text

        checkpoint_text = str(raw_d).strip()
        if not checkpoint_text:
            continue  # empty string

        # Skip if the "checkpoint" text is actually a formula (sanity guard)
        if checkpoint_text.startswith("="):
            continue

        # --- Weight (col M) ---
        raw_m = _cell_value(ws, row, col_m)
        weight = _to_float_weight(raw_m)

        # Skip header-like rows that have text but no weight
        if weight == 0.0:
            logger.debug("Row %d skipped — no weight (text: %s)", row, checkpoint_text[:40])
            continue

        # --- Existing verdict (col L) — warn if already filled ---
        existing_verdict = _cell_value(ws, row, col_l)
        if existing_verdict is not None and str(existing_verdict).strip():
            logger.warning(
                "Checkpoint Q2_2026_%d already has verdict '%s' — will be overwritten on this run.",
                row, existing_verdict
            )

        checkpoint_id = f"Q2_2026_{row}"
        excel_cell      = f"{col_l}{row}"   # e.g. "L12"
        observation_cell = f"{col_q}{row}"  # e.g. "Q12"

        checkpoints.append({
            "checkpoint_id":    checkpoint_id,
            "phase_number":     last_phase_number,
            "phase_name":       last_phase_name,
            "activity":         last_activity,
            "checkpoint_text":  checkpoint_text,
            "weight":           weight,
            "excel_cell":       excel_cell,
            "observation_cell": observation_cell,
            "excel_sheet":      sheet_name,
        })

    wb.close()

    # --- Assertions ---
    actual = len(checkpoints)
    assert actual == expected, (
        f"Expected {expected} checkpoints but got {actual}. "
        "Check config.yaml column mapping — the count must match exactly."
    )

    ids = [c["checkpoint_id"] for c in checkpoints]
    assert len(ids) == len(set(ids)), "Duplicate checkpoint_ids detected — fix row numbering logic."

    # --- Write to disk ---
    out = Path(output_path)
    out.write_text(json.dumps(checkpoints, indent=2, ensure_ascii=False), encoding="utf-8")
    logger.info("Written %d checkpoints to %s", actual, output_path)

    return checkpoints


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys

    cfg_path = "config.yaml"
    with open(cfg_path, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)

    result = extract_checkpoints(cfg)
    print(f"\nSUCCESS — {len(result)} checkpoints extracted.")
    print(f"Sample (first entry):\n{json.dumps(result[0], indent=2)}")
    print(f"Sample (last entry):\n{json.dumps(result[-1], indent=2)}")

    # Quick activity count
    activities = {c["activity"] for c in result}
    print(f"\nDistinct activities: {len(activities)}")
