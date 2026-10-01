"""
tests/test_excel_writer.py
==========================
Unit tests for excel_writer.py.
Uses a copy of the real template to verify write correctness.
"""
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
import pytest
import yaml

sys.path.insert(0, str(Path(__file__).parent.parent))

from src.assessor import CheckpointVerdict
from src.excel_writer import build_write_instructions, write_scorecard
from src.template_reader import extract_checkpoints

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def config():
    with open("config.yaml", encoding="utf-8") as f:
        return yaml.safe_load(f)


@pytest.fixture(scope="module")
def registry(config, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("out") / "registry.json"
    checkpoints = extract_checkpoints(config, output_path=str(tmp))
    return {c["checkpoint_id"]: c for c in checkpoints}


def _make_verdict(cid, verdict, obs="[00:01:00 | PM] Evidence sentence."):
    return CheckpointVerdict(
        checkpoint_id=cid, verdict=verdict, confidence=0.8, observation=obs
    )


# ---------------------------------------------------------------------------
# build_write_instructions tests
# ---------------------------------------------------------------------------

def test_na_excluded_from_instructions(registry):
    verdicts = [
        _make_verdict("Q2_2026_12", "Yes"),
        _make_verdict("Q2_2026_13", "NA"),
        _make_verdict("Q2_2026_14", "No"),
    ]
    instructions = build_write_instructions(verdicts, registry)
    cells = [i[1] for i in instructions]
    assert "L13" not in cells, "NA verdict should not generate an instruction"
    assert "L12" in cells
    assert "L14" in cells


def test_only_two_cells_per_instruction(registry):
    """Each instruction must have 5 fields: sheet, L-cell, verdict, Q-cell, observation."""
    verdicts = [_make_verdict("Q2_2026_12", "Yes")]
    instructions = build_write_instructions(verdicts, registry)
    assert len(instructions) == 1
    assert len(instructions[0]) == 5


def test_verdict_cell_is_column_L(registry):
    verdicts = [_make_verdict("Q2_2026_12", "Yes"), _make_verdict("Q2_2026_48", "Partial")]
    instructions = build_write_instructions(verdicts, registry)
    for _, verdict_cell, *_ in instructions:
        assert verdict_cell.startswith("L"), f"Verdict cell not in column L: {verdict_cell}"


def test_observation_cell_is_column_Q(registry):
    verdicts = [_make_verdict("Q2_2026_12", "Yes")]
    instructions = build_write_instructions(verdicts, registry)
    _, _, _, obs_cell, _ = instructions[0]
    assert obs_cell.startswith("Q"), f"Observation cell not in column Q: {obs_cell}"


def test_no_instructions_for_all_na(registry):
    verdicts = [_make_verdict("Q2_2026_12", "NA"), _make_verdict("Q2_2026_13", "NA")]
    instructions = build_write_instructions(verdicts, registry)
    assert instructions == []


# ---------------------------------------------------------------------------
# write_scorecard tests
# ---------------------------------------------------------------------------

@pytest.fixture
def template_copy(tmp_path):
    src = Path("data/template.xlsx")
    dest = tmp_path / "template_copy.xlsx"
    shutil.copy2(str(src), str(dest))
    return str(dest)


def test_assertion_blocks_na_write(template_copy, registry):
    """Attempting to pass an NA verdict to write_scorecard raises AssertionError."""
    # Manually construct an instruction with NA bypassing Point A
    bad_instruction = ("Q2 2026", "L12", "NA", "Q12", "")
    out = template_copy.replace("template_copy", "out_bad")
    with pytest.raises(AssertionError, match="NA verdict"):
        write_scorecard(template_copy, [bad_instruction], out)


def test_verdict_written_correctly(template_copy, registry):
    """Written verdict string should be readable back from the cell."""
    verdicts = [_make_verdict("Q2_2026_12", "Yes")]
    instructions = build_write_instructions(verdicts, registry)
    out = template_copy.replace("template_copy", "out_verdict")
    write_scorecard(template_copy, instructions, out)

    wb = openpyxl.load_workbook(out, data_only=True)
    ws = wb["Q2 2026"]
    assert ws["L12"].value == "Yes", f"Expected 'Yes', got {ws['L12'].value}"
    wb.close()


def test_observation_written_correctly(template_copy, registry):
    obs = "[00:14:32 | Project Manager] SOW was reviewed and approved."
    verdicts = [_make_verdict("Q2_2026_12", "Yes", obs=obs)]
    instructions = build_write_instructions(verdicts, registry)
    out = template_copy.replace("template_copy", "out_obs")
    write_scorecard(template_copy, instructions, out)

    wb = openpyxl.load_workbook(out, data_only=True)
    ws = wb["Q2 2026"]
    assert ws["Q12"].value == obs, f"Observation mismatch: {ws['Q12'].value!r}"
    wb.close()


def test_na_cell_untouched(template_copy, registry):
    """NA verdict — the L and Q cells must remain completely blank."""
    # Write only non-NA verdicts; do NOT write any instruction for Q2_2026_13
    verdicts = [
        _make_verdict("Q2_2026_12", "Yes"),
        _make_verdict("Q2_2026_13", "NA"),  # should be filtered out
    ]
    instructions = build_write_instructions(verdicts, registry)
    out = template_copy.replace("template_copy", "out_na")
    write_scorecard(template_copy, instructions, out)

    wb = openpyxl.load_workbook(out, data_only=True)
    ws = wb["Q2 2026"]
    l13 = ws["L13"].value
    q13 = ws["Q13"].value
    wb.close()
    assert l13 is None or l13 == "", f"L13 should be untouched, got {l13!r}"
    assert q13 is None or q13 == "", f"Q13 should be untouched, got {q13!r}"


def test_formula_cells_n_not_overwritten(template_copy, registry):
    """N-column formula cells must remain formulas after writing."""
    verdicts = [_make_verdict("Q2_2026_12", "Yes")]
    instructions = build_write_instructions(verdicts, registry)
    out = template_copy.replace("template_copy", "out_formula")
    write_scorecard(template_copy, instructions, out)

    # Load with data_only=False to read formula strings
    wb = openpyxl.load_workbook(out, data_only=False)
    ws = wb["Q2 2026"]
    n12 = ws["N12"].value
    wb.close()
    assert isinstance(n12, str) and n12.startswith("="), (
        f"N12 should be a formula, got {n12!r}"
    )


def test_original_template_unchanged(template_copy, registry):
    """The original template must not be modified."""
    orig_src = Path("data/template.xlsx")
    orig_size = orig_src.stat().st_size

    verdicts = [_make_verdict("Q2_2026_12", "Yes")]
    instructions = build_write_instructions(verdicts, registry)
    out = template_copy.replace("template_copy", "out_orig_check")
    write_scorecard(template_copy, instructions, out)

    # Original template size should not change
    assert orig_src.stat().st_size == orig_size, "Original template was modified!"


def test_in_progress_written_correctly(template_copy, registry):
    """'In Progress' should be written and read back correctly."""
    verdicts = [_make_verdict("Q2_2026_12", "In Progress")]
    instructions = build_write_instructions(verdicts, registry)
    out = template_copy.replace("template_copy", "out_inprogress")
    write_scorecard(template_copy, instructions, out)

    wb = openpyxl.load_workbook(out, data_only=True)
    ws = wb["Q2 2026"]
    val = ws["L12"].value
    wb.close()
    assert val == "In Progress", f"Expected 'In Progress', got {val!r}"
